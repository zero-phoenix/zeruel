#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""tools/docx_utils.py: Utilidades reutilizables de manipulación y limpieza OOXML para Zeruel."""

import os
import re
import zipfile
import docx

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"


def clean_highlights_zip(docx_path: str) -> None:
    """Elimina toda etiqueta w:highlight en document.xml, footnotes, headers, footers."""
    temp_zip = docx_path + ".tmp.zip"
    with zipfile.ZipFile(docx_path, 'r') as zin:
        with zipfile.ZipFile(temp_zip, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.endswith('.xml'):
                    text = data.decode('utf-8', 'replace')
                    text = re.sub(r'<w:highlight\b[^>]*/>', '', text)
                    text = re.sub(r'<w:highlight\b[^>]*>.*?</w:highlight>', '', text)
                    zout.writestr(item, text.encode('utf-8'))
                else:
                    zout.writestr(item, data)
    os.replace(temp_zip, docx_path)


def replace_paragraph_text(p, new_text: str, bold_prefix: str = None) -> None:
    """Reemplaza el texto de un párrafo preservando su estilo y run base."""
    r_font = "Arial Narrow"
    r_size = None
    if p.runs:
        r0 = p.runs[0]
        if r0.font.name:
            r_font = r0.font.name
        if r0.font.size:
            r_size = r0.font.size

    p.text = ""  # vacía runs preservando pPr

    if bold_prefix and new_text.startswith(bold_prefix):
        resto = new_text[len(bold_prefix):]
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.name = r_font
        if r_size:
            r1.font.size = r_size

        r2 = p.add_run(resto)
        r2.bold = False
        r2.font.name = r_font
        if r_size:
            r2.font.size = r_size
    else:
        r = p.add_run(new_text)
        r.font.name = r_font
        if r_size:
            r.font.size = r_size
