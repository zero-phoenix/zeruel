#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""tools/inspect_new_cases.py: Extrae y resume contenido de expedientes y borradores
sin acoplamiento a rutas fijas de usuario."""

import os
import sys
import argparse
from pathlib import Path

try:
    import fitz
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

DEFAULT_CASES_DIR = Path(os.environ.get("ZERUEL_CORRECTIONS_DIR", Path.home() / "Desktop" / "Nuevas correcciones"))


def parse_case(base_dir: Path, folder_name: str):
    base = base_dir / folder_name
    print(f"\n=======================================================")
    print(f"ANÁLISIS DETALLADO DEL CASO: {folder_name}")
    print(f"=======================================================")

    if not base.exists():
        print(f"[WARN] Directorio del caso no existe: {base}")
        return

    # 1. Leer el borrador Word
    doc_path = base / f"ADM {folder_name} R2 borrador.docx"
    if doc_path.exists() and docx:
        doc = docx.Document(str(doc_path))
        print(f"\n--- BORRADOR WORD: {len(doc.paragraphs)} párrafos ---")
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if t:
                if any(k in t.upper() for k in ["EXPEDIENTE", "DENUNCIANTE", "DENUNCIADO", "MATERIA", "RESOLUCIÓN", "HECHOS", "INFRACCIÓN", "ARTÍCULO", "RESUELVE", "PRIMERO", "SEGUNDO", "TERCERO", "CUARTO", "QUINTO", "DÉCIMO"]):
                    print(f"  [P{i:02d}] {t[:140]}")
                elif "consistente en que" in t or "reclamo" in t.lower() or "siniestro" in t.lower():
                    print(f"  [P{i:02d}] {t[:140]}")
    elif not doc_path.exists():
        print(f"No se encontró borrador: {doc_path}")

    # 2. Leer los PDFs
    print(f"\n--- DOCUMENTOS PDF DEL EXPEDIENTE ---")
    for f in sorted(base.iterdir()):
        if f.suffix.lower() == ".pdf" and fitz:
            pdf = fitz.open(str(f))
            full_text = ""
            for page in pdf:
                full_text += page.get_text() + "\n"
            print(f"\n* PDF: {f.name} ({len(pdf)} págs, {len(full_text)} caracteres)")
            limpio = " ".join(full_text.split())
            print(f"  Inicio: {limpio[:350]}")


def main():
    parser = argparse.ArgumentParser(description="Inspecciona expedientes y borradores en un directorio.")
    parser.add_argument("--cases-dir", type=Path, default=DEFAULT_CASES_DIR, help="Directorio base que contiene los casos")
    parser.add_argument("--cases", nargs="*", default=["3017-2026", "3057-2026", "3075-2026"], help="Nombres de las carpetas de casos a procesar")
    args = parser.parse_args()

    if not args.cases_dir.exists():
        print(f"[WARN] Directorio base no encontrado: {args.cases_dir}")
        print("Especifique un directorio válido mediante --cases-dir o la variable ZERUEL_CORRECTIONS_DIR.")
        sys.exit(0)

    for c in args.cases:
        parse_case(args.cases_dir, c)


if __name__ == "__main__":
    main()
