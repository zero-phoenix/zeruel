"""Read local workload sources serially; export only allowlisted aggregate metadata.

Never emits text, filenames, sheet names, identities, case IDs or reversible mappings.
Originals remain read-only. Requires openpyxl and pypdf in the operator's local runtime.
This is evidence extraction, not legal decision-making or universal anonymization.
"""
import argparse
from collections import Counter
from datetime import date, datetime
import json
from pathlib import Path
import re
import unicodedata
import zipfile
import xml.etree.ElementTree as ET

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
FIELDS = {
    'expediente_origen': ('expediente de origen',),
    'expediente': ('expediente',), 'fecha_limite': ('fecha limite', 'f. limite'),
    'presentacion': ('presentacion',), 'dias_habiles_transcurridos': ('d_h', 'dh transcurridos'),
    'grupo': ('grupo', 'sector'), 'asignacion': ('asistente', 'analista', 'encargado', 'responsable'),
    'observacion': ('observacion',), 'denunciante': ('denunciante',),
    'denunciado': ('denunciado', 'denunciada'), 'fecha': ('fecha',),
    'estado': ('estado', 'situacion'), 'materia': ('materia',),
}
SECTIONS = ('antecedentes', 'analisis', 'cuestiones en discusion', 'sumilla',
            'vistos', 'se resuelve', 'resuelve', 'resolucion', 'notificacion', 'anexos')
KINDS = ('admisorio', 'requerimiento', 'inadmisibilidad', 'improcedencia',
         'apelacion', 'notificacion', 'sin_clasificar')

def norm(value):
    return ''.join(c for c in unicodedata.normalize('NFD', str(value).lower())
                   if unicodedata.category(c) != 'Mn')

def labels(value):
    if not isinstance(value, str) or len(value) > 160:
        return []
    n = norm(value)
    return [k for k, words in FIELDS.items() if any(word in n for word in words)]

def case_keys(text):
    # Local-only joins, never returned in the public JSON.
    return set(re.findall(r'\b0*(\d{1,6})\s*-\s*(20\d{2})\s*/\s*(CC1|APEL)', text.upper()))

def workbook(path, alias, report=False):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    result, insurance_ids = [], set()
    try:
        for si, sheet in enumerate(book.worksheets, 1):
            iterator = sheet.iter_rows(values_only=True)
            prefix = []
            for _ in range(20):
                row = next(iterator, None)
                if row is None: break
                prefix.append(row)
            if not prefix: continue
            hi = max(range(len(prefix)), key=lambda i: sum(bool(labels(v)) for v in prefix[i]))
            fields = {str(i + 1): labels(v) or ['campo_no_clasificado']
                      for i, v in enumerate(prefix[hi]) if v is not None}
            # Report headers are row 3; do not mistake data rows for the header.
            if report and len(prefix) >= 3:
                hi = 2
                fields = {str(i + 1): labels(v) or ['campo_no_clasificado']
                          for i, v in enumerate(prefix[hi]) if v is not None}
            field_at = {int(c) - 1: f for c, f in fields.items()}
            group_cols = [c for c, f in field_at.items() if 'grupo' in f]
            id_cols = [c for c, f in field_at.items() if 'expediente' in f]
            days_cols = [c for c, f in field_at.items() if 'dias_habiles_transcurridos' in f]
            rows_all = rows_insurance = 0
            ids = set()
            ages, dates_by_field = [], {}
            unknown_group = 0
            def consume(row):
                nonlocal rows_all, rows_insurance, unknown_group
                if not any(v is not None for v in row): return
                if id_cols and not any(c < len(row) and row[c] is not None for c in id_cols): return
                rows_all += 1
                insured = any(c < len(row) and norm(row[c]).strip() == 'seguros' for c in group_cols)
                if not group_cols:
                    insured = 'seguro' in norm(sheet.title)
                    if not insured: unknown_group += 1
                if not insured: return
                rows_insurance += 1
                for c in id_cols:
                    if c < len(row) and row[c] is not None:
                        key = norm(row[c]).strip()
                        ids.add(key)
                        insurance_ids.update(case_keys(str(row[c])))
                for c in days_cols:
                    if c < len(row) and isinstance(row[c], (int, float)) and not isinstance(row[c], bool):
                        ages.append(row[c])
                for c, value in enumerate(row):
                    if isinstance(value, (date, datetime)):
                        field = '/'.join(field_at.get(c, ['fecha_sin_etiqueta']))
                        dates_by_field.setdefault(field, Counter())[value.strftime('%Y-%m')] += 1
            for row in prefix[hi + 1:]: consume(row)
            for row in iterator: consume(row)
            title = norm(sheet.title + ' ' + str(prefix[0][0] or ''))
            role = ('sin_admitir' if 'sin admitir' in title else 'requerimientos' if 'requer' in title
                    else 'apelaciones' if 'apel' in title else 'base_general')
            result.append({'alias': f'{alias}_hoja_{si}', 'rol': role, 'campos': fields,
                'filas_datos': rows_all, 'filas_seguros': rows_insurance,
                'identificadores_distintos_seguros_en_hoja': len(ids),
                'filas_sin_clasificacion_de_sector': unknown_group,
                'dias_habiles': {'n': len(ages), 'min': min(ages, default=None),
                    'max': max(ages, default=None), 'valores_mayores_20': sum(a > 20 for a in ages)},
                'fechas_por_mes': {k: dict(v) for k, v in dates_by_field.items()}})
    finally:
        book.close()
    return result, insurance_ids

def docx(path):
    with zipfile.ZipFile(path) as archive:
        xml = ET.fromstring(archive.read('word/document.xml'))
        paras = [''.join(t.text or '' for t in p.iter(W + 't')) for p in xml.iter(W + 'p')]
        text = '\n'.join(paras)
        fonts, sizes, alignments, spacing, margins = Counter(), Counter(), Counter(), Counter(), Counter()
        for node in xml.iter(W + 'rFonts'):
            raw = node.get(W + 'ascii', '')
            allowed = ('Arial', 'Calibri', 'Times New Roman', 'Cambria', 'Verdana', 'Tahoma', 'Georgia', 'Courier New')
            fonts[raw if raw in allowed else 'otra_o_desconocida'] += 1
        for node in xml.iter(W + 'sz'):
            raw = node.get(W + 'val', '')
            if raw.isdigit(): sizes[str(int(raw) / 2)] += 1
        for node in xml.iter(W + 'jc'):
            raw = node.get(W + 'val', '')
            alignments[raw if raw in ('left', 'right', 'center', 'both') else 'otro'] += 1
        for node in xml.iter(W + 'spacing'):
            raw = node.get(W + 'line', '')
            if raw.isdigit(): spacing[f'{int(raw)}_twips_{node.get(W+"lineRule", "auto")}'] += 1
        for node in xml.iter(W + 'pgMar'):
            values = [node.get(W + side, '') for side in ('top', 'right', 'bottom', 'left')]
            if all(v.isdigit() for v in values): margins['/'.join(values)] += 1
        header_count = sum(name.startswith('word/header') and name.endswith('.xml') for name in archive.namelist())
        footer_count = sum(name.startswith('word/footer') and name.endswith('.xml') for name in archive.namelist())
        return text, {'paragraphs': len(paras), 'tables': len(list(xml.iter(W + 'tbl'))),
            'headers': header_count, 'footers': footer_count,
            'fonts_direct': dict(fonts), 'size_pt_direct': dict(sizes), 'alignments_direct': dict(alignments),
            'spacing_direct': dict(spacing), 'margins_twips_trbl': dict(margins)}

def classify(text, filename):
    n, fn = norm(text), norm(filename)
    if 'cedula de notificacion' in n[:5000] or re.search(r'\bcn\b|cedula', fn): return 'notificacion'
    if 'recurso de apelacion' in n or '/apel' in n or 'apelac' in fn: return 'apelacion'
    operative = n[n.rfind('resuelve'):] if 'resuelve' in n else n
    if re.search(r'declarar.{0,80}inadmisible', operative, re.S): return 'inadmisibilidad'
    if re.search(r'declarar.{0,80}improcedente', operative, re.S): return 'improcedencia'
    if re.search(r'admitir.{0,40}tramite', operative, re.S): return 'admisorio'
    if 'requerir' in operative or 'subsane' in operative or 'subsanacion' in operative: return 'requerimiento'
    return 'sin_clasificar'

def analyze(documents, reports):
    excel, insurance_ids = [], set()
    for folder, is_report in ((documents, False), (reports, True)):
        for path in sorted(folder.rglob('*.xlsx')):
            rows, ids = workbook(path, f'libro_{len(excel)+1}', is_report)
            excel.append({'origen': 'reporte_semanal' if is_report else 'base_general', 'hojas': rows})
            insurance_ids.update(ids)
    cases_by_kind = {kind: set() for kind in KINDS}
    formats, kinds, scopes, sections, numeric_periods = Counter(), Counter(), Counter(), Counter(), Counter()
    metadata, failures, scans, all_ids = [], Counter(), 0, set()
    for index, path in enumerate(sorted(p for p in documents.rglob('*') if p.suffix.lower() in ('.docx', '.pdf')), 1):
        ext = path.suffix.lower()
        formats[ext] += 1
        try:
            if ext == '.docx':
                text, layout = docx(path)
            else:
                from pypdf import PdfReader
                reader = PdfReader(path)
                text = '\n'.join(page.extract_text() or '' for page in reader.pages)
                layout = {'pages': len(reader.pages), 'extractable_text': bool(text.strip())}
                if not text.strip(): scans += 1
            kind, n, ids = classify(text, path.name), norm(text), case_keys(text)
            scope = ('seguros_por_enlace_local' if ids & insurance_ids else
                     'seguros_por_indicios_textuales' if re.search(r'aseguradora|poliza|contrato de seguro|compan[ií]a de seguros|siniestro', n)
                     else 'sector_no_confirmado')
            kinds[kind] += 1; scopes[scope] += 1; all_ids.update(ids); cases_by_kind[kind].update(ids)
            detected = [section for section in SECTIONS if section in n]
            sections.update(detected)
            # Export frequencies, never source sentences or case-specific dates.
            for number in re.findall(r'\b(\d{1,3})\s*(?:\([^)]{1,20}\)\s*)?dias?\s+habiles', n):
                numeric_periods[number] += 1
            metadata.append({'alias': f'documento_{index:03}', 'tipo_archivo': ext, 'familia_provisional': kind,
                'alcance_provisional': scope, 'secciones_detectadas': detected, 'formato': layout})
            del text
        except Exception as exc:
            failures[type(exc).__name__] += 1
            metadata.append({'alias': f'documento_{index:03}', 'tipo_archivo': ext, 'lectura': 'fallo',
                             'error_tipo': type(exc).__name__})
    notifications = cases_by_kind['notificacion']
    resolutions = set().union(*(ids for kind, ids in cases_by_kind.items() if kind != 'notificacion'))
    return {'schema': 1, 'privacy': 'aggregate_and_allowlisted_structure_only',
        'scope': 'seguros; legal classifications provisional until review', 'workbooks': excel,
        'documents': {'file_counts': dict(formats), 'read': len(metadata)-sum(failures.values()),
            'read_failures': dict(failures), 'pdf_without_extractable_text': scans,
            'provisional_families_all_sources': dict(kinds), 'provisional_scope': dict(scopes),
            'section_frequency': dict(sections), 'business_day_mentions_not_rules': dict(numeric_periods),
            'distinct_case_keys_local_all_sources': len(all_ids),
            'cases_with_resolution_and_notification': len(notifications & resolutions),
            'cases_with_resolution_no_notification_found': len(resolutions - notifications),
            'cases_with_notification_no_resolution_found': len(notifications - resolutions)},
        'document_structure': metadata,
        'limitations': ['No raw text, source names, case IDs, identity values or reversible mappings exported.',
            'Word direct formatting is not a full effective-style or visual fidelity verification.',
            'A case join is not proof of one notification for each resolution or of effective delivery.',
            'An unknown sector is never assigned to insurance automatically.',
            'Counts describe source snapshots, not current live assignments or completed case counts.']}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--documents', type=Path, required=True)
    parser.add_argument('--reports', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if not args.documents.is_dir() or not args.reports.is_dir():
        parser.error('Expected two existing local source directories')
    result = analyze(args.documents, args.reports)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    # Do not print original paths, exceptions or extracted source contents.
    print(json.dumps({'written': True, 'documents': result['documents'],
        'workbooks': result['workbooks']}, ensure_ascii=False))

if __name__ == '__main__':
    main()
